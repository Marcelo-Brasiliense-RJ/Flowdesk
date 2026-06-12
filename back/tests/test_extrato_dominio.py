"""Testes unitários do template extrato-dominio (funções puras, dados sintéticos)."""
import sys
import importlib.util
import datetime as dt
from pathlib import Path

BACK = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BACK))


def _load_template_module(tmp_path, monkeypatch):
    from app.routers.templates import TEMPLATES
    from app.services.storage import _SDK_SOURCE
    (tmp_path / "flowdesk_sdk.py").write_text(_SDK_SOURCE, encoding="utf-8")
    code = TEMPLATES["extrato-dominio"]["code"]
    f = tmp_path / "processar.py"
    f.write_text(code, encoding="utf-8")
    run = tmp_path / "run"
    run.mkdir(exist_ok=True)
    (tmp_path / "input.json").write_text('{"_somente_definicoes": true}', encoding="utf-8")
    monkeypatch.setenv("FLOWDESK_INPUT", str(tmp_path / "input.json"))
    monkeypatch.setenv("FLOWDESK_OUTPUT", str(tmp_path / "output.json"))
    monkeypatch.setenv("FLOWDESK_OUTPUT_DIR", str(tmp_path / "out"))
    monkeypatch.setenv("FLOWDESK_RUN_DIR", str(run))
    monkeypatch.syspath_prepend(str(tmp_path))
    spec = importlib.util.spec_from_file_location("processar_t", f)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)  # _somente_definicoes=True faz main() não rodar
    return mod


def test_padrao_remove_numeros(tmp_path, monkeypatch):
    m = _load_template_module(tmp_path, monkeypatch)
    assert m.padrao_de("CIELO VDA DEBITO MASTER 8207205") == "CIELO VDA DEBITO MASTER"
    assert m.padrao_de("REM: JOINT BILLION BRAZIL 01/04") == "REM: JOINT BILLION BRAZIL"
    assert m.padrao_de("TARIFA BANCARIA 300326") == "TARIFA BANCARIA"


def test_gerar_xlsx_contrato_dominio(tmp_path, monkeypatch):
    m = _load_template_module(tmp_path, monkeypatch)
    lanc = [
        {"data": dt.date(2026, 4, 1), "historico": "VENDA CARTAO", "credito": 100.5, "debito": None},
        {"data": dt.date(2026, 4, 1), "historico": "TARIFA", "credito": None, "debito": 8.14},
        {"data": dt.date(2026, 4, 2), "historico": "PIX REC", "credito": 50.0, "debito": None},
    ]
    mapa = {"VENDA CARTAO": "412", "TARIFA": "777", "PIX REC": "413"}
    out = tmp_path / "saida.xlsx"
    m.gerar_xlsx(lanc, mapa, conta_banco="9", caminho=out)
    import openpyxl
    ws = openpyxl.load_workbook(out)["Planilha1"]
    headers = [c.value for c in ws[1]]
    assert headers == ["Data", "Cód. Conta Debito", "Cód. Conta Credito", "Valor",
                       "Cód. Histórico", "Complemento Histórico", "Inicia Lote",
                       "Código Matriz/Filial", "Centro de Custo Débito",
                       "Centro de Custo Crédito"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    # entrada no banco: débito banco 9 / crédito contrapartida; Inicia Lote=1 só na 1a linha do dia
    assert rows[0][1] == "9" and rows[0][2] == "412" and rows[0][6] == 1
    assert rows[1][1] == "777" and rows[1][2] == "9" and rows[1][6] is None
    assert rows[2][6] == 1  # novo dia 02/04 reinicia lote
    assert rows[0][4] is None  # Cód. Histórico sempre vazio
    assert abs(rows[0][3] - 100.5) < 0.001  # valor numérico


def test_classificar_com_regras_e_ignorar(tmp_path, monkeypatch):
    m = _load_template_module(tmp_path, monkeypatch)
    lanc = [
        {"data": dt.date(2026, 4, 1), "historico": "TARIFA BANCARIA 1", "credito": None, "debito": 5.0},
        {"data": dt.date(2026, 4, 1), "historico": "COISA NOVA 99", "credito": 10.0, "debito": None},
    ]
    regras = {"TARIFA BANCARIA": "777"}
    grupos = m.agrupar(lanc, regras)
    por_padrao = {g["padrao"]: g for g in grupos}
    assert por_padrao["TARIFA BANCARIA"]["conta"] == "777"
    assert por_padrao["TARIFA BANCARIA"]["origem"] == "regra"
    assert por_padrao["COISA NOVA"]["conta"] is None
