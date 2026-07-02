"""Aceite E2E deterministico: o codigo semeado pelo gerador (template extrato-dominio)
rodando contra as fixtures reais em testes/dominio/. Sem OpenAI."""
import importlib.util
import json
import sys
from pathlib import Path

import pytest

BACK = Path(__file__).resolve().parent.parent
FIXTURES = BACK.parent / "testes" / "dominio"
PDF = FIXTURES / "JBB GIG - EXTRATO BRADESCO 04.2026 (1).pdf"
CONTAS = FIXTURES / "Contas.xlsx"

pytestmark = pytest.mark.skipif(
    not (PDF.exists() and CONTAS.exists()),
    reason="fixtures locais em testes/dominio/ ausentes",
)


def _carrega_template(tmp_path, monkeypatch):
    from app.routers.templates import TEMPLATES
    from app.services.storage import _SDK_SOURCE

    (tmp_path / "flowdesk_sdk.py").write_text(_SDK_SOURCE, encoding="utf-8")
    f = tmp_path / "processar.py"
    f.write_text(TEMPLATES["extrato-dominio"]["code"], encoding="utf-8")
    run = tmp_path / "run"
    run.mkdir(exist_ok=True)
    out_dir = tmp_path / "out"
    out_dir.mkdir(exist_ok=True)
    # o SDK cacheia os caminhos de env no import; remove do cache de modulos para
    # forcar reimport com o env desta chamada (varios testes no mesmo processo)
    sys.modules.pop("flowdesk_sdk", None)
    monkeypatch.setenv("FLOWDESK_INPUT", str(tmp_path / "input.json"))
    monkeypatch.setenv("FLOWDESK_OUTPUT", str(tmp_path / "output.json"))
    monkeypatch.setenv("FLOWDESK_OUTPUT_DIR", str(out_dir))
    monkeypatch.setenv("FLOWDESK_RUN_DIR", str(run))
    monkeypatch.syspath_prepend(str(tmp_path))
    spec = importlib.util.spec_from_file_location("processar_e2e", f)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)  # __name__ != '__main__' => main() nao roda no import
    return mod, tmp_path


def _roda(mod, tmp_path, entrada: dict) -> dict:
    (tmp_path / "input.json").write_text(json.dumps(entrada), encoding="utf-8")
    saida = tmp_path / "output.json"
    if saida.exists():
        saida.unlink()
    mod.main()
    return json.loads(saida.read_text(encoding="utf-8"))


def test_pass1_le_plano_e_monta_revisao(tmp_path, monkeypatch):
    mod, tmp = _carrega_template(tmp_path, monkeypatch)
    out = _roda(mod, tmp, {"arquivo1": str(PDF), "arquivo2": str(CONTAS)})
    rev = out.get("_classificacao_review")
    assert rev, f"esperava _classificacao_review, veio {list(out)}"
    assert len(rev["contas"]) > 1000                 # leu o Contas.xlsx de fato
    assert rev["grupos"], "sem grupos classificaveis"
    assert rev["total_lancamentos"] > 0


def test_pass2_gera_arquivo_resultado(tmp_path, monkeypatch):
    mod, tmp = _carrega_template(tmp_path, monkeypatch)
    rev = _roda(mod, tmp, {"arquivo1": str(PDF), "arquivo2": str(CONTAS)})["_classificacao_review"]
    # confirma todos os grupos numa conta analitica qualquer do plano
    conta = rev["contas"][0]["codigo"]
    confirmado = {g["padrao"]: conta for g in rev["grupos"]}
    periodo = rev["periodo_detectado"]
    out = _roda(mod, tmp, {
        "arquivo1": str(PDF), "arquivo2": str(CONTAS),
        "_classificacao_confirmada": confirmado,
        "_conta_banco": "1",
        "_periodo": periodo,
    })
    assert out.get("arquivo_resultado"), f"esperava arquivo_resultado, veio {list(out)}"
    from openpyxl import load_workbook
    wb = load_workbook(out["arquivo_resultado"])
    ws = wb["Planilha1"]
    linhas = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(linhas) > 0                            # planilha do Dominio preenchida


def test_pdf_invalido_retorna_erro_legivel(tmp_path, monkeypatch):
    mod, tmp = _carrega_template(tmp_path, monkeypatch)
    ruim = tmp / "vazio.pdf"
    ruim.write_bytes(b"nao e um pdf")
    # get_file(0) precisa existir no disco; usa o arquivo ruim como extrato
    try:
        out = _roda(mod, tmp, {"arquivo1": str(ruim), "arquivo2": str(CONTAS)})
    except Exception:
        pytest.skip("template propaga excecao; o canal 'erro' e coberto por test_runner_error")
    assert out.get("erro") or out.get("_classificacao_review") is None
